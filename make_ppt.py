# -*- coding: utf-8 -*-
"""
MY AI LIFE - PowerPoint Presentation Generator
16:9 와이드스크린, 모던 다크/사이버 테마 슬라이드 자동 생성 스크립트
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path="My_AI_Life_Presentation.pptx"):
    prs = Presentation()
    # 16:9 Widescreen (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # 빈 슬라이드

    # 색상 팔레트 정의
    BG_COLOR = RGBColor(15, 23, 42)        # #0f172a (다크 슬레이트 네이비)
    CARD_BG = RGBColor(30, 41, 59)         # #1e293b (딥 네이비 카드)
    CARD_BG_ALT = RGBColor(24, 32, 47)     # 더 짙은 카드
    BORDER_COLOR = RGBColor(51, 65, 85)    # #334155
    
    TEXT_WHITE = RGBColor(248, 250, 252)   # #f8fafc
    TEXT_MUTED = RGBColor(148, 163, 184)   # #94a3b8
    TEXT_DIM = RGBColor(100, 116, 139)     # #64748b
    
    ACCENT_BLUE = RGBColor(56, 189, 248)   # #38bdf8 (사이버 블루)
    ACCENT_CYAN = RGBColor(34, 211, 238)   # #22d3ee (시안)
    ACCENT_GREEN = RGBColor(52, 211, 153)  # #34d399 (민트/에메랄드)
    ACCENT_PURPLE = RGBColor(168, 85, 247) # #a855f7 (바이올렛)
    ACCENT_AMBER = RGBColor(251, 191, 36)  # #fbbf24 (골드 옐로우)
    ACCENT_ROSE = RGBColor(244, 63, 94)    # #f43f5e (로즈 레드)

    FONT_MAIN = "Malgun Gothic"
    FONT_TITLE = "Malgun Gothic"

    def apply_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, tag_color=ACCENT_BLUE):
        # 상단 태그
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.name = FONT_MAIN
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = tag_color

        # 메인 제목
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = FONT_TITLE
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # ==========================================
    # SLIDE 1: 표지 (Title Slide)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_background(s1)

    # 장식용 백그라운드 카드
    center_card = add_card(s1, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1), CARD_BG, ACCENT_BLUE)

    # 표지 텍스트 박스
    t_box = s1.shapes.add_textbox(Inches(2.0), Inches(1.8), Inches(9.333), Inches(3.8))
    tf = t_box.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "INTERACTIVE LIFE SIMULATION & OSMU ENGINE"
    p0.font.name = FONT_MAIN
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "MY AI LIFE"
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "나의 AI 인생 시뮬레이터 : 프로젝트 핵심 요약"
    p2.font.name = FONT_MAIN
    p2.font.size = Pt(20)
    p2.font.color.rgb = ACCENT_CYAN
    p2.space_after = Pt(20)

    p3 = tf.add_paragraph()
    p3.text = "AI의 극효율 알고리즘과 인간의 가치관이 빚어내는 인생의 갈림길,\n그리고 나만의 선택으로 완성되는 웹툰 & 숏폼 콘텐츠 생성 엔진"
    p3.font.name = FONT_MAIN
    p3.font.size = Pt(13)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_after = Pt(28)

    p4 = tf.add_paragraph()
    p4.text = "개발 스택: Vanilla JS (ES6 Modules) • Zero-Dependency • Responsive Web"
    p4.font.name = FONT_MAIN
    p4.font.size = Pt(11)
    p4.font.color.rgb = TEXT_DIM

    notes1 = s1.notes_slide.notes_text_frame
    notes1.text = "안녕하십니까, 데이터와 가치관의 갈림길을 시뮬레이션하는 인터랙티브 프로젝트 'My AI Life'의 발표를 시작하겠습니다."

    # ==========================================
    # SLIDE 2: 기획 배경 및 문제 정의
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_background(s2)
    add_header(s2, "01. Background & Problem", "기획 배경 및 문제 정의 : AI 시대의 삶과 선택")

    # 좌측 카드 (Problem)
    add_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.2), CARD_BG, ACCENT_ROSE)
    tb_prob = s2.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(3.6))
    tf_p = tb_prob.text_frame
    tf_p.word_wrap = True

    p = tf_p.paragraphs[0]
    p.text = "PROBLEM : 알고리즘과 극효율의 역설"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE
    p.space_after = Pt(12)

    bullets_p = [
        "데이터와 AI가 진로·투자·직업의 '최적 경로'를 제시하는 시대",
        "수치와 소득 중심의 극효율 추구로 인한 번아웃과 자아실현 결핍",
        "질문: 'AI가 계산한 정답이 과연 나에게도 가장 행복한 삶일까?'",
        "단기적 선택이 20~30년 뒤 인생에 미칠 인과관계를 체감하기 어려움"
    ]
    for b in bullets_p:
        pb = tf_p.add_paragraph()
        pb.text = "• " + b
        pb.font.name = FONT_MAIN
        pb.font.size = Pt(12)
        pb.font.color.rgb = TEXT_MUTED
        pb.space_after = Pt(8)

    # 우측 카드 (Solution)
    add_card(s2, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.2), CARD_BG, ACCENT_GREEN)
    tb_sol = s2.shapes.add_textbox(Inches(7.2), Inches(2.1), Inches(5.0), Inches(3.6))
    tf_s = tb_sol.text_frame
    tf_s.word_wrap = True

    p = tf_s.paragraphs[0]
    p.text = "SOLUTION : 인터랙티브 라이프 시뮬레이터"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(12)

    bullets_s = [
        "10대부터 50대까지 현실 공공 통계(팩트 카드) 기반 딜레마 체험",
        "AI의 극효율 권고와 나의 가치관을 비교하며 독창적 삶의 궤적 발견",
        "청년기 결정이 중장년기에 발동되는 지연 나비효과 큐(Queue) 시각화",
        "평가나 서열이 아닌, '자신만의 선택'이 지닌 고유한 가치 회고"
    ]
    for b in bullets_s:
        pb = tf_s.add_paragraph()
        pb.text = "• " + b
        pb.font.name = FONT_MAIN
        pb.font.size = Pt(12)
        pb.font.color.rgb = TEXT_MUTED
        pb.space_after = Pt(8)

    # 하단 콜아웃
    add_card(s2, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.8), CARD_BG_ALT, ACCENT_CYAN)
    tb_call = s2.shapes.add_textbox(Inches(1.1), Inches(6.35), Inches(11.1), Inches(0.5))
    p = tb_call.text_frame.paragraphs[0]
    p.text = "💡 핵심 가치: 효율과 행복은 결코 같은 단어가 아니다 — 선택의 무게와 주체성을 되찾는 여정"
    p.font.name = FONT_MAIN
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    # ==========================================
    # SLIDE 3: 3대 핵심 차별점 (Core Features)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_background(s3)
    add_header(s3, "02. Core Features", "핵심 차별점 : 3대 인터랙티브 시스템")

    cols = [
        {
            "title": "1. 6대 라이프 HUD",
            "tag": "다차원 상태 지표",
            "color": ACCENT_BLUE,
            "items": [
                "💰 자산 (Money) - 시드머니/부동산",
                "💼 커리어 (Career) - 전문성과 직위",
                "📚 역량 (Skill) - 기술 및 리스킬링",
                "🤝 관계 (Relation) - 가족 및 유대감",
                "⏳ 여유 시간 (Time) - 번아웃 방지",
                "🌱 삶의 만족도 (Life) - 자아실현",
                "➔ 6대 스탯의 실시간 트레이드오프"
            ]
        },
        {
            "title": "2. 진화형 AI 페르소나",
            "tag": "3단계 파트너",
            "color": ACCENT_PURPLE,
            "items": [
                "🤖 [10~20대] The Optimizer",
                "  : 데이터 기반의 극효율 경로 제시",
                "⚡ [30~40대] The Challenger",
                "  : 기회비용과 내면 가치관 질문",
                "🪞 [50대~] The Reflector",
                "  : 삶의 궤적과 인간성을 비추는 거울",
                "➔ 연령대별 차별화된 실시간 코멘트"
            ]
        },
        {
            "title": "3. 팩트 카드 & 나비효과",
            "tag": "인과관계 엔진",
            "color": ACCENT_AMBER,
            "items": [
                "📊 대한민국 공공 통계 팩트 카드",
                "🔍 2단계 숙고 (결정 번복 추적)",
                "🦋 지연 나비효과 큐(Queue)",
                "  - 17세 창작 열정 ➔ 24세 스튜디오",
                "  - 34세 영끌 ➔ 45세 금리 압박",
                "  - 45세 리스킬 ➔ 56세 신산업 리더",
                "➔ 과거 선택이 미래에 실시간 영향"
            ]
        }
    ]

    card_w = Inches(3.64)
    card_gap = Inches(0.39)
    start_x = Inches(0.8)

    for i, c in enumerate(cols):
        cx = start_x + i * (card_w + card_gap)
        add_card(s3, cx, Inches(1.8), card_w, Inches(5.1), CARD_BG, c["color"])
        tb = s3.shapes.add_textbox(cx + Inches(0.25), Inches(2.05), card_w - Inches(0.5), Inches(4.6))
        tf_c = tb.text_frame
        tf_c.word_wrap = True

        p = tf_c.paragraphs[0]
        p.text = c["tag"].upper()
        p.font.name = FONT_MAIN
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = c["color"]

        p_t = tf_c.add_paragraph()
        p_t.text = c["title"]
        p_t.font.name = FONT_TITLE
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_after = Pt(12)

        for item in c["items"]:
            pi = tf_c.add_paragraph()
            pi.text = item
            pi.font.name = FONT_MAIN
            pi.font.size = Pt(11)
            pi.font.color.rgb = TEXT_MUTED
            pi.space_after = Pt(4)

    # ==========================================
    # SLIDE 4: 게임 루프 & 2단계 숙고 메커니즘
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_background(s4)
    add_header(s4, "03. Game Mechanics", "게임 메커니즘 : 2단계 숙고 및 나비효과 루프")

    steps = [
        {
            "step": "STEP 01",
            "title": "상황 마주하기",
            "color": ACCENT_BLUE,
            "desc": "• 연령별 현실적 딜레마 발생\n• AI 극효율 알고리즘 권고 및 정량적 근거 확인\n• 3가지 분기 선택지 탐색"
        },
        {
            "step": "STEP 02",
            "title": "팩트 카드 확인",
            "color": ACCENT_CYAN,
            "desc": "• 1차 직관적 선택 진행\n• 실제 정부/학술 통계 데이터 모달 팝업 오픈\n• 숨겨진 기회비용과 리스크 수치 인지"
        },
        {
            "step": "STEP 03",
            "title": "2차 숙고 & 확정",
            "color": ACCENT_AMBER,
            "desc": "• 팩트를 바탕으로 결정 유지 or 번복\n• AI 권고 수용/거절 여부 추적\n• 가치관 축 가중치 누적 반영"
        },
        {
            "step": "STEP 04",
            "title": "나비효과 & 피드백",
            "color": ACCENT_GREEN,
            "desc": "• 6대 Life Status 즉각 갱신\n• 미래 특정 연령에 터질 나비효과 큐 등록\n• AI 페르소나의 실시간 평가 코멘터리"
        }
    ]

    s_w = Inches(2.68)
    s_gap = Inches(0.32)
    s_start = Inches(0.8)

    for i, st in enumerate(steps):
        sx = s_start + i * (s_w + s_gap)
        add_card(s4, sx, Inches(1.8), s_w, Inches(3.6), CARD_BG, st["color"])
        tb = s4.shapes.add_textbox(sx + Inches(0.2), Inches(2.0), s_w - Inches(0.4), Inches(3.2))
        tf_s = tb.text_frame
        tf_s.word_wrap = True

        p = tf_s.paragraphs[0]
        p.text = st["step"]
        p.font.name = FONT_MAIN
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = st["color"]

        pt = tf_s.add_paragraph()
        pt.text = st["title"]
        pt.font.name = FONT_TITLE
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_WHITE
        pt.space_after = Pt(10)

        pd = tf_s.add_paragraph()
        pd.text = st["desc"]
        pd.font.name = FONT_MAIN
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_MUTED

    # 하단 설명 박스
    add_card(s4, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.4), CARD_BG_ALT, BORDER_COLOR)
    tb_inf = s4.shapes.add_textbox(Inches(1.1), Inches(5.75), Inches(11.1), Inches(1.1))
    tf_inf = tb_inf.text_frame
    tf_inf.word_wrap = True
    p = tf_inf.paragraphs[0]
    p.text = "🎯 숙고(Deliberation) 추적 엔진의 의의"
    p.font.name = FONT_TITLE
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    p.space_after = Pt(4)

    p2 = tf_inf.add_paragraph()
    p2.text = "단순 클릭 게임이 아닌 '팩트 통계를 확인한 후 결정을 바꿨는가?'(Fact Change Rate)와 'AI 권고를 거절하고 내 가치를 지켰는가?'(AI Rejection Rate)를 통계로 누적하여 최종 리포트에 인간적 고뇌의 궤적으로 환원합니다."
    p2.font.name = FONT_MAIN
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 5: 생애 주기별 시나리오 여정 (10s ~ 50s)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_background(s5)
    add_header(s5, "04. Scenarios Journey", "시나리오 구성 : 10대부터 50대까지의 5대 갈림길")

    scenarios_data = [
        ("17세 (10대)", "진로의 첫 갈림길: 꿈인가 안정인가", "공학/상경 진학 집중\n(취업 안정성 +42%)", "안정 진학 vs 학업·창작 병행 vs 창작 올인"),
        ("26세 (20대)", "첫 커리어: 스타트업 vs 대기업 계약직", "대기업 계약직 선택\n(시드머니 32% 우세)", "대기업 계약직 vs AI 스타트업 vs 프리랜서"),
        ("34세 (30대)", "자산 형성: 영끌 내집마련 vs 유연성", "아파트 자가 매수\n(부동산 레버리지 최적화)", "자가 매수 vs 유연 투자/자유 vs 워라밸/가족"),
        ("45세 (40대)", "전환기 생존: AI 격변 속 이직 vs 창업", "야간 AI 리스킬링\n(대체 리스크 헷징)", "직장 수성 vs AI 리스킬링 vs 독립 창업"),
        ("56세 (50대)", "인생 2막: 자녀 지원 vs 노후 자금", "연금 전액 사수\n(노후 빈곤 차단)", "연금 전액 사수 vs 절충형 지원 vs 전폭 지원")
    ]

    row_h = Inches(0.9)
    table_top = Inches(1.8)
    
    for idx, (age, title, ai_rec, choices) in enumerate(scenarios_data):
        ry = table_top + idx * Inches(1.0)
        card = add_card(s5, Inches(0.8), ry, Inches(11.7), Inches(0.9), CARD_BG, BORDER_COLOR)
        
        # 연령 뱃지
        tb_age = s5.shapes.add_textbox(Inches(1.0), ry + Inches(0.2), Inches(1.6), Inches(0.5))
        p = tb_age.text_frame.paragraphs[0]
        p.text = age
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

        # 시나리오 제목
        tb_t = s5.shapes.add_textbox(Inches(2.7), ry + Inches(0.2), Inches(3.4), Inches(0.5))
        p = tb_t.text_frame.paragraphs[0]
        p.text = title
        p.font.name = FONT_MAIN
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

        # AI 추천
        tb_ai = s5.shapes.add_textbox(Inches(6.2), ry + Inches(0.12), Inches(2.6), Inches(0.65))
        tf_a = tb_ai.text_frame
        tf_a.word_wrap = True
        p = tf_a.paragraphs[0]
        p.text = "AI 극효율: " + ai_rec.split('\n')[0]
        p.font.name = FONT_MAIN
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = ACCENT_AMBER
        if len(ai_rec.split('\n')) > 1:
            p2 = tf_a.add_paragraph()
            p2.text = ai_rec.split('\n')[1]
            p2.font.size = Pt(9.5)
            p2.font.color.rgb = TEXT_MUTED

        # 선택지 분기
        tb_ch = s5.shapes.add_textbox(Inches(8.9), ry + Inches(0.2), Inches(3.4), Inches(0.5))
        p = tb_ch.text_frame.paragraphs[0]
        p.text = choices
        p.font.name = FONT_MAIN
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 6: 듀얼 엔딩 & 평행우주 리포트
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_background(s6)
    add_header(s6, "05. Dual Ending", "엔딩 시스템 : MY LIFE vs AI PARALLEL LIFE")

    # 좌측 (MY LIFE)
    add_card(s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.3), CARD_BG, ACCENT_GREEN)
    tb_my = s6.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(3.7))
    tf_my = tb_my.text_frame
    tf_my.word_wrap = True

    p = tf_my.paragraphs[0]
    p.text = "【 MY LIFE (나의 선택) 】"
    p.font.name = FONT_TITLE
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(10)

    bullets_my = [
        "나의 가치관과 고유한 감정이 온전히 반영된 여정",
        "자산은 적을지라도 풍요로운 관계와 여유 시간 확보",
        "지연된 나비효과가 빚어낸 개성 넘치는 독창적 서사",
        "연령별 선택 궤적을 잇는 '당신이 살아온 또 하나의 인생' 헌정사 제공",
        "AI 거부율(%) 및 팩트 기반 결정 번복률(%) 통계 분석"
    ]
    for b in bullets_my:
        pb = tf_my.add_paragraph()
        pb.text = "✔ " + b
        pb.font.name = FONT_MAIN
        pb.font.size = Pt(11.5)
        pb.font.color.rgb = TEXT_MUTED
        pb.space_after = Pt(6)

    # 우측 (AI PARALLEL LIFE)
    add_card(s6, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.3), CARD_BG, ACCENT_PURPLE)
    tb_ai = s6.shapes.add_textbox(Inches(7.2), Inches(2.1), Inches(5.0), Inches(3.7))
    tf_ai = tb_ai.text_frame
    tf_ai.word_wrap = True

    p = tf_ai.paragraphs[0]
    p.text = "【 AI LIFE (100% 극효율 경로) 】"
    p.font.name = FONT_TITLE
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE
    p.space_after = Pt(10)

    bullets_ai = [
        "ParallelLifeEngine이 백그라운드에서 완전 자동 시뮬레이션",
        "수치적으로 극대화된 자산(Money)과 커리어(Career) 점수",
        "극효율을 위해 과감히 희생된 인간관계, 여유 시간, 삶의 만족도",
        "듀얼 레이더 차트를 통한 나의 결과와의 직관적 오버레이 비교",
        "'AI의 정답대로만 살았을 때의 나'를 마주하는 철학적 충격"
    ]
    for b in bullets_ai:
        pb = tf_ai.add_paragraph()
        pb.text = "🤖 " + b
        pb.font.name = FONT_MAIN
        pb.font.size = Pt(11.5)
        pb.font.color.rgb = TEXT_MUTED
        pb.space_after = Pt(6)

    # 하단 콜아웃
    add_card(s6, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7), CARD_BG_ALT, ACCENT_CYAN)
    tb_c = s6.shapes.add_textbox(Inches(1.1), Inches(6.42), Inches(11.1), Inches(0.45))
    p = tb_c.text_frame.paragraphs[0]
    p.text = "✨ 서열화와 승패 판정을 배제하고, '나의 선택이 만들어낸 고유한 삶의 궤적'을 회고하는 힐링 리포트"
    p.font.name = FONT_MAIN
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    # ==========================================
    # SLIDE 7: One Life Source : OSMU 콘텐츠 자동 생성
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_background(s7)
    add_header(s7, "06. OSMU Content Engine", "확장성 : 시뮬레이션에서 웹툰·숏폼 콘텐츠로 (One Life Source)")

    # 좌측 (웹툰 콘티)
    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), CARD_BG, ACCENT_BLUE)
    tb_wt = s7.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.5))
    tf_wt = tb_wt.text_frame
    tf_wt.word_wrap = True

    p = tf_wt.paragraphs[0]
    p.text = "🎨 인스타툰/웹툰 4~5컷 콘티 스크립트"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(10)

    wt_cuts = [
        "1컷 (17세): 갈림길 앞 주인공 & AI 홀로그램 대사",
        "  - AI: '상경계열 진학이 취업 안정성 +42%입니다'",
        "  - 주인공: '나는 창작의 길을 걷겠어!'",
        "2컷 (26세): 스타트업 야생 vs 대기업 고뇌 묘사",
        "3컷 (34세): 팩트 카드 오버레이와 내면 가치 충돌",
        "4컷 (56세): 수십 년 뒤 발동된 나비효과 결말 씬",
        "➔ 웹툰 작가/일러스트레이터가 즉시 작화 가능한 표준 콘티 포맷 자동 렌더링"
    ]
    for w in wt_cuts:
        pw = tf_wt.add_paragraph()
        pw.text = ("• " if not w.startswith("  ") and not w.startswith("➔") else "") + w
        pw.font.name = FONT_MAIN
        pw.font.size = Pt(10.5)
        pw.font.color.rgb = TEXT_WHITE if w.startswith("➔") else TEXT_MUTED
        pw.font.bold = True if w.startswith("➔") else False
        pw.space_after = Pt(3)

    # 우측 (숏폼 & Google Vids 연동)
    add_card(s7, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.1), CARD_BG, ACCENT_PURPLE)
    tb_sh = s7.shapes.add_textbox(Inches(7.2), Inches(2.1), Inches(5.0), Inches(4.5))
    tf_sh = tb_sh.text_frame
    tf_sh.word_wrap = True

    p = tf_sh.paragraphs[0]
    p.text = "🎬 60초 숏폼 & Google Vids AI 연동"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE
    p.space_after = Pt(10)

    sh_lines = [
        "[00:00 - 00:05] HOOK: 'AI 말만 들으면 인생 성공할까?'",
        "[00:05 - 00:45] DILEMMA: AI 권고 vs 나의 실제 선택 팝업",
        "[00:45 - 00:55] BUTTERFLY: 20년 뒤 나비효과 반전",
        "[00:55 - 01:00] OUTRO: MY LIFE vs AI LIFE 듀얼 엔딩",
        "✨ Google Vids (Help me create a video) 원클릭 연동",
        "  - 9:16 Vertical Shorts 규격에 맞춘 씬별 프롬프트 자동 생성",
        "  - [프롬프트 복사] & [Vids 열기]로 1분 만에 AI 보이스/영상 완성"
    ]
    for s in sh_lines:
        ps = tf_sh.add_paragraph()
        ps.text = ("• " if not s.startswith("  ") and not s.startswith("✨") else "") + s
        ps.font.name = FONT_MAIN
        ps.font.size = Pt(10.5)
        ps.font.color.rgb = ACCENT_CYAN if s.startswith("✨") else (TEXT_WHITE if s.startswith("  -") else TEXT_MUTED)
        ps.font.bold = True if s.startswith("✨") or s.startswith("[") else False
        ps.space_after = Pt(3)

    # ==========================================
    # SLIDE 8: 시스템 아키텍처 및 기술 스택
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_background(s8)
    add_header(s8, "07. Tech Stack & Architecture", "시스템 아키텍처 : 모듈형 설계 및 Zero-Dependency")

    # 좌측 아키텍처 박스
    add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), CARD_BG, BORDER_COLOR)
    tb_arch = s8.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.5))
    tf_ar = tb_arch.text_frame
    tf_ar.word_wrap = True

    p = tf_ar.paragraphs[0]
    p.text = "📁 모듈형 소프트웨어 구조"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(10)

    arch_text = (
        "my_ai_life/\n"
        "├── js/\n"
        "│   ├── app.js               # 진입점 & 오케스트레이션\n"
        "│   ├── engine/\n"
        "│   │   ├── LifeEngine.js    # 코어 시뮬레이션/나비효과\n"
        "│   │   └── ParallelLifeEngine.js # AI 평행우주 엔진\n"
        "│   ├── ui/\n"
        "│   │   ├── HUD.js           # 6대 상태 게이지\n"
        "│   │   ├── AICompanion.js   # AI 페르소나/말풍선\n"
        "│   │   ├── FactCardModal.js # 팩트 카드 2차 숙고\n"
        "│   │   └── ReportView.js    # 듀얼 차트/OSMU 탭\n"
        "│   └── data/\n"
        "│       ├── constants.js     # 초기 상태 및 페르소나\n"
        "│       └── scenarios.js     # 10s~50s 팩트 데이터셋\n"
        "└── css/main.css             # 글래스모피즘 & 반응형 UI"
    )
    pa = tf_ar.add_paragraph()
    pa.text = arch_text
    pa.font.name = "Consolas"
    pa.font.size = Pt(9.5)
    pa.font.color.rgb = TEXT_MUTED

    # 우측 기술적 특장점
    add_card(s8, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.1), CARD_BG, ACCENT_BLUE)
    tb_tech = s8.shapes.add_textbox(Inches(7.2), Inches(2.1), Inches(5.0), Inches(4.5))
    tf_tc = tb_tech.text_frame
    tf_tc.word_wrap = True

    p = tf_tc.paragraphs[0]
    p.text = "⚡ 기술적 특장점 및 엔지니어링"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(12)

    tech_points = [
        ("Zero-Dependency 경량 웹 앱", "무거운 빌드 툴(Webpack/Vite)이나 외부 라이브러리 없이 브라우저에서 즉시 구동 (Live Server 등 완벽 지원)"),
        ("순수 ES6+ 모듈 구조", "Engine(로직), UI(렌더링), Data(시나리오)가 완벽히 분리되어 유지보수와 기능 확장이 매우 용이"),
        ("반응형 모던 다크 UI & 글래스모피즘", "CSS Grid & Flexbox 기반 반응형 레이아웃, 반투명 블러 효과 및 다이내믹 게이지 애니메이션"),
        ("확장 가능한 데이터 모델", "새로운 연령대 시나리오나 추가 라이프 지표를 손쉽게 추가할 수 있는 유연한 데이터 스키마")
    ]
    for title, desc in tech_points:
        pt = tf_tc.add_paragraph()
        pt.text = "• " + title
        pt.font.name = FONT_MAIN
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_WHITE
        
        pd = tf_tc.add_paragraph()
        pd.text = "  " + desc
        pd.font.name = FONT_MAIN
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(8)

    # ==========================================
    # SLIDE 9: 기대 효과 및 향후 발전 로드맵
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_background(s9)
    add_header(s9, "08. Roadmap & Vision", "로드맵 및 비전 : 인터랙티브 에듀테크 & 미디어 플랫폼")

    roadmaps = [
        {
            "phase": "PHASE 01 (v1.0)",
            "status": "현재 완료",
            "color": ACCENT_GREEN,
            "title": "코어 시스템 및 MVP 완성",
            "items": [
                "• 10대~50대 5대 생애주기 시나리오",
                "• 6대 라이프 지표 HUD & 나비효과 큐",
                "• 3단계 진화형 AI 페르소나",
                "• AI 평행우주 듀얼 엔딩 리포트",
                "• 웹툰 콘티 & 숏폼 대본 자동 생성"
            ]
        },
        {
            "phase": "PHASE 02 (v1.5)",
            "status": "단기 확장 계획",
            "color": ACCENT_CYAN,
            "title": "LLM 동적 연동 & 고도화",
            "items": [
                "• LLM(Gemini API) 실시간 연동",
                "• 사용자 자유 입력 맞춤형 시나리오 생성",
                "• AI 파트너와의 양방향 실시간 대화/설득",
                "• 청소년/대학생 진로 및 금융 교육 커리큘럼 연계 B2B 툴 개발"
            ]
        },
        {
            "phase": "PHASE 03 (v2.0)",
            "status": "중장기 비전",
            "color": ACCENT_PURPLE,
            "title": "생성형 AI 멀티미디어 파이프라인",
            "items": [
                "• 이미지/비디오 생성 AI 파이프라인 연동",
                "• 인스타툰 컷 이미지 자동 렌더링",
                "• 숏폼 영상/보이스 원클릭 퍼블리싱",
                "• 글로벌 다국어(EN/JP) 시나리오 확장",
                "• 종합 라이프 컨설팅 플랫폼 도약"
            ]
        }
    ]

    for i, rm in enumerate(roadmaps):
        rx = start_x + i * (card_w + card_gap)
        add_card(s9, rx, Inches(1.8), card_w, Inches(5.1), CARD_BG, rm["color"])
        tb = s9.shapes.add_textbox(rx + Inches(0.25), Inches(2.05), card_w - Inches(0.5), Inches(4.6))
        tf_r = tb.text_frame
        tf_r.word_wrap = True

        p = tf_r.paragraphs[0]
        p.text = rm["phase"] + " [" + rm["status"] + "]"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = rm["color"]

        pt = tf_r.add_paragraph()
        pt.text = rm["title"]
        pt.font.name = FONT_TITLE
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_WHITE
        pt.space_after = Pt(12)

        for it in rm["items"]:
            pi = tf_r.add_paragraph()
            pi.text = it
            pi.font.name = FONT_MAIN
            pi.font.size = Pt(11)
            pi.font.color.rgb = TEXT_MUTED
            pi.space_after = Pt(6)

    # 저장
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_presentation()
